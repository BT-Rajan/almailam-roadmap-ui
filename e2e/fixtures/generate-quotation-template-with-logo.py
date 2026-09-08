"""Regenerates e2e/fixtures/quotation-template-with-logo.docx.

A minimal but realistic admin-authored Quotation template: a letterhead
paragraph that is nothing but an embedded PNG logo, a second paragraph
that mixes a logo image with static text (the kind of letterhead line a
real company template tends to have), a couple of plain paragraphs left
for the admin to map merge fields onto, and a two-column table for the
line-items merge field.

Exists so document-template-mapping.spec.ts has a real .docx -- with a
real embedded image -- to upload through the actual UI, instead of
faking the upload. Run `python3 generate-quotation-template-with-logo.py`
from this directory to regenerate the fixture (needs `pip install
python-docx pillow`).
"""

import io
from pathlib import Path

from docx import Document
from PIL import Image

HERE = Path(__file__).resolve().parent

logo_bytes = io.BytesIO()
Image.new("RGB", (240, 90), color=(178, 34, 34)).save(logo_bytes, format="PNG")
logo_path = HERE / "_tmp_logo.png"
logo_path.write_bytes(logo_bytes.getvalue())

doc = Document()

# Image-only letterhead paragraph -- extract_layout reports this as an
# empty-text block; apply_mapping must never destroy it just because
# the admin's mapping payload includes it unmodified.
doc.add_paragraph().add_run().add_picture(str(logo_path))

# Mixed text + image paragraph -- a realistic "logo beside the company
# name" letterhead line.
p = doc.add_paragraph()
p.add_run("Al Mailam Trading LLC ")
p.add_run().add_picture(str(logo_path))
p.add_run(" - Est. 1995")

# Left blank for the admin to place merge fields onto via Map Fields.
doc.add_paragraph()
doc.add_paragraph()

table = doc.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "Description"
table.rows[0].cells[1].text = "Amount"
# Row 1 left blank for the admin to mark as the repeating line-items row.

doc.save(str(HERE / "quotation-template-with-logo.docx"))
logo_path.unlink()
print("wrote", HERE / "quotation-template-with-logo.docx")
