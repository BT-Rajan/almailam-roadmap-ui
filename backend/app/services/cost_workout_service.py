"""Cost workout document for a Quotation or Contract -- the priced
breakdown (line items, subtotal, discount, total, amount in words, plus
the monthly Supervision fees for reference), generated straight from the
record as .docx and .pdf.

This replaced merging an admin-uploaded, field-mapped Word template for
these two document types: staff no longer pick/maintain a template or
frame a quotation/contract letter -- the cost workout is the document.
(Payment Plan still uses document_template_service's template merge.)
"""

import html
import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from sqlalchemy.orm import Session
from weasyprint import HTML

from app.core.exceptions import ValidationAppError
from app.core.file_storage import resolve_path
from app.core.number_to_words import amount_to_words
from app.models.client import Client
from app.models.contract import Contract
from app.models.project import Project
from app.models.quotation import Quotation
from app.models.user import User
from app.services import company_service
from app.services.pdf_render import FONT_PATH

LANGUAGES = ("English", "Arabic")

_LABELS = {
    "English": {
        "title": "Cost Workout",
        "quotation": "Quotation",
        "contract": "Contract",
        "client": "Client",
        "project": "Project",
        "project_no": "Project No.",
        "site_address": "Project/Site Address",
        "issue_date": "Issue Date",
        "valid_until": "Valid Until",
        "expiry_date": "Expiry Date",
        "revision": "Revision",
        "prepared_by": "Prepared By",
        "no": "#",
        "description": "Description",
        "qty": "Qty",
        "unit_price": "Unit Price",
        "amount": "Amount",
        "subtotal": "Subtotal",
        "discount": "Discount",
        "total": "Total",
        "contract_value": "Contract Value",
        "in_words": "Amount in words",
        "supervision": "Supervision (billed monthly, not included in the total above)",
        "activity": "Activity",
        "monthly_rate": "Monthly Rate",
        "period": "Period",
        "to": "to",
    },
    "Arabic": {
        "title": "تفصيل التكلفة",
        "quotation": "عرض سعر",
        "contract": "عقد",
        "client": "العميل",
        "project": "المشروع",
        "project_no": "رقم المشروع",
        "site_address": "عنوان المشروع/الموقع",
        "issue_date": "تاريخ الإصدار",
        "valid_until": "صالح حتى",
        "expiry_date": "تاريخ الانتهاء",
        "revision": "المراجعة",
        "prepared_by": "أعده",
        "no": "#",
        "description": "الوصف",
        "qty": "الكمية",
        "unit_price": "سعر الوحدة",
        "amount": "المبلغ",
        "subtotal": "المجموع الفرعي",
        "discount": "الخصم",
        "total": "الإجمالي",
        "contract_value": "قيمة العقد",
        "in_words": "المبلغ كتابةً",
        "supervision": "الإشراف (يُحتسب شهرياً، غير مشمول في الإجمالي أعلاه)",
        "activity": "النشاط",
        "monthly_rate": "الأجر الشهري",
        "period": "الفترة",
        "to": "إلى",
    },
}


@dataclass
class _Line:
    description: str
    quantity: Decimal
    unit_price: Decimal

    @property
    def amount(self) -> Decimal:
        return self.quantity * self.unit_price


@dataclass
class _Supervision:
    activity: str
    monthly_rate: Decimal
    start: date
    end: date


@dataclass
class _Workout:
    language: str
    kind: str  # "quotation" | "contract"
    number: str
    company_name: str
    logo_path: Path | None
    # (label key, value) pairs shown in the header block, in order.
    details: list[tuple[str, str]]
    currency: str
    lines: list[_Line]
    discount: Decimal
    total: Decimal
    total_label_key: str
    supervision: list[_Supervision] = field(default_factory=list)

    @property
    def labels(self) -> dict[str, str]:
        return _LABELS[self.language]

    @property
    def subtotal(self) -> Decimal:
        return sum((line.amount for line in self.lines), Decimal("0"))

    @property
    def heading(self) -> str:
        return f"{self.labels['title']} — {self.labels[self.kind]} {self.number}"


def _dec(value) -> Decimal:
    return Decimal(str(value or 0))


def _money(value: Decimal) -> str:
    return f"{value:,.2f}"


def _qty(value: Decimal) -> str:
    return f"{float(value):g}"


def _fmt_date(value: date | None) -> str:
    return value.strftime("%d %B %Y") if value else ""


def _check_language(language: str) -> str:
    if language not in LANGUAGES:
        raise ValidationAppError(f"Unsupported document language '{language}'.")
    return language


def _company(db: Session, language: str | None) -> tuple[str, Path | None, str]:
    settings = company_service.get_settings(db)
    logo = resolve_path(settings.logo_storage_key) if settings.logo_storage_key else None
    if logo is not None and not logo.is_file():
        logo = None
    resolved_language = _check_language(language) if language is not None else settings.default_language
    if resolved_language not in LANGUAGES:
        resolved_language = "English"
    return settings.company_name, logo, resolved_language


def _client_name(client: Client | None) -> str:
    # Same "M/s." convention as document_template_service._client_display_name.
    if not client or not client.company_name:
        return ""
    return f"M/s. {client.company_name}"


def _user_name(db: Session, user_id: int | None) -> str:
    if user_id is None:
        return ""
    user = db.query(User).filter(User.id == user_id).first()
    return user.full_name if user else ""


def _project_and_client(db: Session, project_id: int) -> tuple[Project | None, Client | None]:
    project = db.query(Project).filter(Project.id == project_id).first()
    client = db.query(Client).filter(Client.id == project.client_id).first() if project else None
    return project, client


def _supervision(db: Session, project: Project | None) -> list[_Supervision]:
    if project is None:
        return []
    from app.services.project_service.queries import get_selected_supervision_activities

    return [
        _Supervision(a.activity_name, _dec(a.monthly_rate), a.start_date, a.end_date)
        for a in get_selected_supervision_activities(db, project.id)
    ]


def _project_details(project: Project | None, client: Client | None) -> list[tuple[str, str]]:
    return [
        ("client", _client_name(client)),
        ("project", project.project_name if project else ""),
        ("project_no", project.project_no if project else ""),
        ("site_address", (project.site_address or "") if project else ""),
    ]


def build_quotation_workout(db: Session, quotation: Quotation, language: str | None) -> _Workout:
    from app.services import quotation_service

    company_name, logo, language = _company(db, language)
    project, client = _project_and_client(db, quotation.project_id)
    lines = [
        _Line(item.description, _dec(item.quantity), _dec(item.unit_price))
        for item in quotation_service.get_line_items(db, quotation.id)
    ]
    return _Workout(
        language=language,
        kind="quotation",
        number=quotation.quotation_no,
        company_name=company_name,
        logo_path=logo,
        details=[
            *_project_details(project, client),
            ("issue_date", _fmt_date(quotation.issue_date)),
            ("valid_until", _fmt_date(quotation.validity)),
            ("revision", quotation.revision),
            ("prepared_by", _user_name(db, quotation.prepared_by)),
        ],
        currency=quotation.currency,
        lines=lines,
        discount=_dec(quotation.discount_amount),
        total=_dec(quotation.amount),
        total_label_key="total",
        supervision=_supervision(db, project),
    )


def build_contract_workout(db: Session, contract: Contract, language: str | None) -> _Workout:
    """A contract is priced off the approved quotation it was created from
    (contract_service requires contract_value to match that quotation's
    amount), so its cost workout reuses that quotation's line items."""
    from app.services import quotation_service

    company_name, logo, language = _company(db, language)
    project, client = _project_and_client(db, contract.project_id)
    lines: list[_Line] = []
    discount = Decimal("0")
    if contract.quotation_id is not None:
        quotation = db.query(Quotation).filter(Quotation.id == contract.quotation_id).first()
        if quotation is not None:
            lines = [
                _Line(item.description, _dec(item.quantity), _dec(item.unit_price))
                for item in quotation_service.get_line_items(db, quotation.id)
            ]
            discount = _dec(quotation.discount_amount)
    contract_value = _dec(contract.contract_value)
    if not lines:
        # Older contract with no linked quotation: a single line for its value.
        lines = [_Line(_LABELS[language]["contract_value"], Decimal("1"), contract_value)]
    return _Workout(
        language=language,
        kind="contract",
        number=contract.contract_no,
        company_name=company_name,
        logo_path=logo,
        details=[
            *_project_details(project, client),
            ("issue_date", _fmt_date(contract.issue_date)),
            ("expiry_date", _fmt_date(contract.expiry_date)),
            ("revision", contract.revision),
            ("prepared_by", _user_name(db, contract.prepared_by)),
        ],
        currency=contract.currency,
        lines=lines,
        discount=discount,
        total=contract_value,
        total_label_key="contract_value",
        supervision=_supervision(db, project),
    )


# -- DOCX ---------------------------------------------------------------------


def _set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _set_rtl(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    p_pr.append(bidi)


def _cell_text(cell, text: str, *, bold: bool = False, align=None, rtl: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    if rtl:
        _set_rtl(paragraph)
        run.font.rtl = True
    if align is not None:
        paragraph.alignment = align


def render_docx(workout: _Workout) -> bytes:
    rtl = workout.language == "Arabic"
    labels = workout.labels
    start = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    end = WD_ALIGN_PARAGRAPH.LEFT if rtl else WD_ALIGN_PARAGRAPH.RIGHT

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(18))
    base = doc.styles["Normal"]
    base.font.name = "Arial"
    base.font.size = Pt(10)

    def para(text: str, *, bold: bool = False, size: int | None = None, align=start):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if rtl:
            _set_rtl(p)
            run.font.rtl = True
        p.alignment = align
        return p

    if workout.logo_path is not None:
        try:
            doc.add_picture(str(workout.logo_path), width=Mm(35))
            doc.paragraphs[-1].alignment = start
        except Exception:
            pass  # An unreadable logo shouldn't block the document.
    para(workout.company_name, bold=True, size=14)
    para(workout.heading, bold=True, size=16)

    details = [(k, v) for k, v in workout.details if v]
    info = doc.add_table(rows=len(details), cols=2)
    info.style = "Table Grid"
    for row, (key, value) in zip(info.rows, details):
        _set_cell_shading(row.cells[0], "F1F3F5")
        _cell_text(row.cells[0], labels[key], bold=True, align=start, rtl=rtl)
        _cell_text(row.cells[1], value, align=start, rtl=rtl)
    doc.add_paragraph()

    headers = [labels["no"], labels["description"], labels["qty"],
               f"{labels['unit_price']} ({workout.currency})", f"{labels['amount']} ({workout.currency})"]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, headers):
        _set_cell_shading(cell, "DDE3EA")
        _cell_text(cell, text, bold=True, align=start, rtl=rtl)
    for index, line in enumerate(workout.lines, start=1):
        cells = table.add_row().cells
        _cell_text(cells[0], str(index), align=start, rtl=rtl)
        _cell_text(cells[1], line.description, align=start, rtl=rtl)
        _cell_text(cells[2], _qty(line.quantity), align=end, rtl=rtl)
        _cell_text(cells[3], _money(line.unit_price), align=end, rtl=rtl)
        _cell_text(cells[4], _money(line.amount), align=end, rtl=rtl)

    totals = [(labels["subtotal"], workout.subtotal, False)]
    if workout.discount:
        totals.append((labels["discount"], -workout.discount, False))
    totals.append((labels[workout.total_label_key], workout.total, True))
    for label, value, bold in totals:
        cells = table.add_row().cells
        merged = cells[0].merge(cells[3])
        _cell_text(merged, label, bold=bold, align=end, rtl=rtl)
        _cell_text(cells[4], f"{_money(value)} {workout.currency}", bold=bold, align=end, rtl=rtl)

    doc.add_paragraph()
    para(f"{labels['in_words']}: {amount_to_words(workout.total, workout.currency)}", bold=True)

    if workout.supervision:
        doc.add_paragraph()
        para(labels["supervision"], bold=True, size=11)
        sup = doc.add_table(rows=1, cols=3)
        sup.style = "Table Grid"
        for cell, text in zip(sup.rows[0].cells, [labels["activity"], f"{labels['monthly_rate']} ({workout.currency})", labels["period"]]):
            _set_cell_shading(cell, "DDE3EA")
            _cell_text(cell, text, bold=True, align=start, rtl=rtl)
        for item in workout.supervision:
            cells = sup.add_row().cells
            _cell_text(cells[0], item.activity, align=start, rtl=rtl)
            _cell_text(cells[1], _money(item.monthly_rate), align=end, rtl=rtl)
            _cell_text(cells[2], f"{_fmt_date(item.start)} {labels['to']} {_fmt_date(item.end)}", align=start, rtl=rtl)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# -- PDF ----------------------------------------------------------------------


def _e(value: object) -> str:
    return html.escape(str(value), quote=True)


def render_pdf(workout: _Workout) -> bytes:
    rtl = workout.language == "Arabic"
    labels = workout.labels
    direction = "rtl" if rtl else "ltr"
    start, end = ("right", "left") if rtl else ("left", "right")
    font_stack = "'NotoNaskhArabic', sans-serif" if rtl else "'Helvetica Neue', Arial, 'NotoNaskhArabic', sans-serif"

    logo_html = ""
    if workout.logo_path is not None:
        logo_html = f'<img class="logo" src="{_e(workout.logo_path.resolve().as_uri())}" alt="">'

    detail_rows = "".join(
        f"<tr><th>{_e(labels[key])}</th><td>{_e(value)}</td></tr>" for key, value in workout.details if value
    )
    line_rows = "".join(
        f"<tr><td>{i}</td><td>{_e(line.description)}</td><td class='num'>{_e(_qty(line.quantity))}</td>"
        f"<td class='num'>{_e(_money(line.unit_price))}</td><td class='num'>{_e(_money(line.amount))}</td></tr>"
        for i, line in enumerate(workout.lines, start=1)
    )
    totals = [(labels["subtotal"], workout.subtotal, "")]
    if workout.discount:
        totals.append((labels["discount"], -workout.discount, ""))
    totals.append((labels[workout.total_label_key], workout.total, "grand"))
    total_rows = "".join(
        f"<tr class='{cls}'><td colspan='4' class='num'>{_e(label)}</td>"
        f"<td class='num'>{_e(_money(value))} {_e(workout.currency)}</td></tr>"
        for label, value, cls in totals
    )

    supervision_html = ""
    if workout.supervision:
        rows = "".join(
            f"<tr><td>{_e(s.activity)}</td><td class='num'>{_e(_money(s.monthly_rate))}</td>"
            f"<td>{_e(_fmt_date(s.start))} {_e(labels['to'])} {_e(_fmt_date(s.end))}</td></tr>"
            for s in workout.supervision
        )
        supervision_html = f"""
<h3>{_e(labels['supervision'])}</h3>
<table class="lines">
  <thead><tr><th>{_e(labels['activity'])}</th><th>{_e(labels['monthly_rate'])} ({_e(workout.currency)})</th><th>{_e(labels['period'])}</th></tr></thead>
  <tbody>{rows}</tbody>
</table>"""

    html_doc = f"""<!DOCTYPE html>
<html dir="{direction}" lang="{'ar' if rtl else 'en'}">
<head>
<meta charset="utf-8">
<style>
  @font-face {{ font-family: 'NotoNaskhArabic'; src: url('{FONT_PATH.as_uri()}'); }}
  @page {{ size: A4 portrait; margin: 18mm; }}
  body {{ font-family: {font_stack}; direction: {direction}; text-align: {start}; font-size: 10pt; color: #1a1a2e; line-height: 1.5; }}
  .logo {{ max-height: 22mm; max-width: 45mm; display: block; margin-bottom: 4mm; }}
  .company {{ font-size: 14pt; font-weight: bold; margin: 0; }}
  h1 {{ font-size: 16pt; margin: 2mm 0 5mm; }}
  h3 {{ font-size: 11pt; margin: 7mm 0 2mm; }}
  table {{ width: 100%; border-collapse: collapse; margin-bottom: 4mm; }}
  th, td {{ border: 1px solid #b8c0c8; padding: 1.6mm 2.4mm; vertical-align: top; text-align: {start}; }}
  .details th {{ width: 32%; background: #f1f3f5; }}
  .lines thead th {{ background: #dde3ea; }}
  .num {{ text-align: {end}; white-space: nowrap; }}
  .grand td {{ font-weight: bold; background: #f1f3f5; }}
  .words {{ font-weight: bold; margin-top: 3mm; }}
</style>
</head>
<body>
{logo_html}
<p class="company">{_e(workout.company_name)}</p>
<h1>{_e(workout.heading)}</h1>
<table class="details">{detail_rows}</table>
<table class="lines">
  <thead><tr>
    <th>{_e(labels['no'])}</th><th>{_e(labels['description'])}</th><th class="num">{_e(labels['qty'])}</th>
    <th class="num">{_e(labels['unit_price'])} ({_e(workout.currency)})</th><th class="num">{_e(labels['amount'])} ({_e(workout.currency)})</th>
  </tr></thead>
  <tbody>{line_rows}{total_rows}</tbody>
</table>
<p class="words">{_e(labels['in_words'])}: {_e(amount_to_words(workout.total, workout.currency))}</p>
{supervision_html}
</body>
</html>"""
    # base_url=None plus only file:// URIs we built ourselves (font, logo)
    # -- no user-supplied text can introduce a fetchable URL, since every
    # value above is HTML-escaped.
    return HTML(string=html_doc).write_pdf()
