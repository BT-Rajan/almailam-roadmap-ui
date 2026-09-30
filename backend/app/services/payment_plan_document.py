"""Built-in Payment Plan document -- used whenever no admin-uploaded
Payment Plan template exists for the requested language (see
document_template_service._resolve_payment_plan_template), so Download /
Print / Email work out of the box instead of failing with "No default
... template is configured". An uploaded template, once there, still
takes precedence.

Renders the exact same context dict render_payment_plan_document builds
for a template merge (project/client header, one summary row per billing
stream, one flat instalment schedule), in the same visual style as the
Quotation/Contract cost workout (cost_workout_service).
"""

import io
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from weasyprint import HTML

from app.services.cost_workout_service import _cell_text, _e, _set_cell_shading, _set_rtl
from app.services.pdf_render import FONT_PATH

_LABELS = {
    "English": {
        "title": "Payment Plan",
        "client": "Client",
        "project": "Project",
        "project_no": "Project No.",
        "site_address": "Project/Site Address",
        "issue_date": "Issue Date",
        "summary": "Summary",
        "stream": "Service",
        "status": "Status",
        "amount": "Amount",
        "payment_mode": "Payment Mode",
        "agreement_date": "Agreement Date",
        "start_date": "Start Date",
        "schedule": "Payment Schedule",
        "no": "#",
        "description": "Description",
        "due_date": "Due Date",
        "amount_due": "Amount Due",
        "total": "Total",
        "empty": "No payment plan has been created for this project yet.",
        "streams": {"Design": "Design & Permit", "Supervision": "Supervision"},
    },
    "Arabic": {
        "title": "خطة الدفع",
        "client": "العميل",
        "project": "المشروع",
        "project_no": "رقم المشروع",
        "site_address": "عنوان المشروع/الموقع",
        "issue_date": "تاريخ الإصدار",
        "summary": "الملخص",
        "stream": "الخدمة",
        "status": "الحالة",
        "amount": "المبلغ",
        "payment_mode": "طريقة الدفع",
        "agreement_date": "تاريخ الاتفاقية",
        "start_date": "تاريخ البدء",
        "schedule": "جدول الدفعات",
        "no": "#",
        "description": "الوصف",
        "due_date": "تاريخ الاستحقاق",
        "amount_due": "المبلغ المستحق",
        "total": "الإجمالي",
        "empty": "لم يتم إنشاء خطة دفع لهذا المشروع بعد.",
        "streams": {"Design": "التصميم والتراخيص", "Supervision": "الإشراف"},
    },
}


def _labels(language: str) -> dict:
    return _LABELS.get(language, _LABELS["English"])


def _money(value: str) -> str:
    return f"{Decimal(value or '0'):,.2f}"


def _details(context: dict, labels: dict) -> list[tuple[str, str]]:
    rows = [
        (labels["client"], context.get("client_name", "")),
        (labels["project"], context.get("project_name", "")),
        (labels["project_no"], context.get("project_no", "")),
        (labels["site_address"], context.get("project_address", "")),
        (labels["issue_date"], context.get("issue_date", "")),
    ]
    return [(label, value) for label, value in rows if value]


def _totals_by_currency(schedule: list[dict]) -> list[tuple[str, Decimal]]:
    totals: dict[str, Decimal] = {}
    for row in schedule:
        totals[row["currency"]] = totals.get(row["currency"], Decimal("0")) + Decimal(row["amount_due"] or "0")
    return list(totals.items())


def render_docx(context: dict, language: str, company_name: str, logo_path: Path | None) -> bytes:
    labels = _labels(language)
    rtl = language == "Arabic"
    start = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    end = WD_ALIGN_PARAGRAPH.LEFT if rtl else WD_ALIGN_PARAGRAPH.RIGHT
    stream_name = labels["streams"]

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(18))
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    def para(text: str, *, bold: bool = False, size: int | None = None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if rtl:
            _set_rtl(p)
            run.font.rtl = True
        p.alignment = start

    def table(headers: list[str], rows: list[list[tuple[str, bool]]]):
        # rows: list of [(text, is_numeric)] cells
        tbl = doc.add_table(rows=1, cols=len(headers))
        tbl.style = "Table Grid"
        for cell, text in zip(tbl.rows[0].cells, headers):
            _set_cell_shading(cell, "DDE3EA")
            _cell_text(cell, text, bold=True, align=start, rtl=rtl)
        for row in rows:
            cells = tbl.add_row().cells
            for cell, (text, numeric) in zip(cells, row):
                _cell_text(cell, text, align=end if numeric else start, rtl=rtl)
        return tbl

    if logo_path is not None:
        try:
            doc.add_picture(str(logo_path), width=Mm(35))
            doc.paragraphs[-1].alignment = start
        except Exception:
            pass  # An unreadable logo shouldn't block the document.
    para(company_name, bold=True, size=14)
    para(f"{labels['title']} — {context.get('project_no', '')}", bold=True, size=16)

    details = _details(context, labels)
    info = doc.add_table(rows=len(details), cols=2)
    info.style = "Table Grid"
    for row, (label, value) in zip(info.rows, details):
        _set_cell_shading(row.cells[0], "F1F3F5")
        _cell_text(row.cells[0], label, bold=True, align=start, rtl=rtl)
        _cell_text(row.cells[1], value, align=start, rtl=rtl)
    doc.add_paragraph()

    streams, schedule = context.get("streams", []), context.get("schedule", [])
    if not streams:
        para(labels["empty"])
    else:
        para(labels["summary"], bold=True, size=11)
        table(
            [labels["stream"], labels["status"], labels["amount"], labels["payment_mode"], labels["agreement_date"], labels["start_date"]],
            [[
                (stream_name.get(s["stream"], s["stream"]), False), (s["status"], False),
                (f"{_money(s['amount'])} {s['currency']}", True), (s["payment_mode"], False),
                (s["agreement_date"], False), (s["start_date"], False),
            ] for s in streams],
        )
        doc.add_paragraph()
        para(labels["schedule"], bold=True, size=11)
        tbl = table(
            [labels["no"], labels["stream"], labels["description"], labels["due_date"], labels["amount_due"]],
            [[
                (row["sequence_number"], False), (stream_name.get(row["stream"], row["stream"]), False),
                (row["description"], False), (row["due_date"], False),
                (f"{_money(row['amount_due'])} {row['currency']}", True),
            ] for row in schedule],
        )
        for currency, total in _totals_by_currency(schedule):
            cells = tbl.add_row().cells
            merged = cells[0].merge(cells[3])
            _cell_text(merged, labels["total"], bold=True, align=end, rtl=rtl)
            _cell_text(cells[4], f"{total:,.2f} {currency}", bold=True, align=end, rtl=rtl)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_pdf(context: dict, language: str, company_name: str, logo_path: Path | None) -> bytes:
    labels = _labels(language)
    rtl = language == "Arabic"
    direction = "rtl" if rtl else "ltr"
    start, end = ("right", "left") if rtl else ("left", "right")
    font_stack = "'NotoNaskhArabic', sans-serif" if rtl else "'Helvetica Neue', Arial, 'NotoNaskhArabic', sans-serif"
    stream_name = labels["streams"]

    logo_html = ""
    if logo_path is not None and logo_path.is_file():
        logo_html = f'<img class="logo" src="{_e(logo_path.resolve().as_uri())}" alt="">'
    detail_rows = "".join(f"<tr><th>{_e(label)}</th><td>{_e(value)}</td></tr>" for label, value in _details(context, labels))

    streams, schedule = context.get("streams", []), context.get("schedule", [])
    if not streams:
        body = f"<p>{_e(labels['empty'])}</p>"
    else:
        summary_rows = "".join(
            f"<tr><td>{_e(stream_name.get(s['stream'], s['stream']))}</td><td>{_e(s['status'])}</td>"
            f"<td class='num'>{_e(_money(s['amount']))} {_e(s['currency'])}</td><td>{_e(s['payment_mode'])}</td>"
            f"<td>{_e(s['agreement_date'])}</td><td>{_e(s['start_date'])}</td></tr>"
            for s in streams
        )
        schedule_rows = "".join(
            f"<tr><td>{_e(r['sequence_number'])}</td><td>{_e(stream_name.get(r['stream'], r['stream']))}</td>"
            f"<td>{_e(r['description'])}</td><td>{_e(r['due_date'])}</td>"
            f"<td class='num'>{_e(_money(r['amount_due']))} {_e(r['currency'])}</td></tr>"
            for r in schedule
        )
        total_rows = "".join(
            f"<tr class='grand'><td colspan='4' class='num'>{_e(labels['total'])}</td>"
            f"<td class='num'>{_e(f'{total:,.2f}')} {_e(currency)}</td></tr>"
            for currency, total in _totals_by_currency(schedule)
        )
        body = f"""
<h3>{_e(labels['summary'])}</h3>
<table class="lines">
  <thead><tr><th>{_e(labels['stream'])}</th><th>{_e(labels['status'])}</th><th class="num">{_e(labels['amount'])}</th>
  <th>{_e(labels['payment_mode'])}</th><th>{_e(labels['agreement_date'])}</th><th>{_e(labels['start_date'])}</th></tr></thead>
  <tbody>{summary_rows}</tbody>
</table>
<h3>{_e(labels['schedule'])}</h3>
<table class="lines">
  <thead><tr><th>{_e(labels['no'])}</th><th>{_e(labels['stream'])}</th><th>{_e(labels['description'])}</th>
  <th>{_e(labels['due_date'])}</th><th class="num">{_e(labels['amount_due'])}</th></tr></thead>
  <tbody>{schedule_rows}{total_rows}</tbody>
</table>"""

    html_doc = f"""<!DOCTYPE html>
<html dir="{direction}" lang="{'ar' if rtl else 'en'}">
<head>
<meta charset="utf-8">
<style>
  @font-face {{ font-family: 'NotoNaskhArabic'; src: url('{FONT_PATH.as_uri()}'); }}
  @page {{ size: A4 portrait; margin: 18mm; }}
  body {{ font-family: {font_stack}; direction: {direction}; text-align: {start}; font-size: 10pt; color: #1a1a2e; line-height: 1.5; }}
  .logo {{ max-height: 22mm; max-width: 45mm; display: block; margin-bottom: 4mm; }}
  .company {{ font-size: 14pt; font-weight: bold; margin: 0; }}
  h1 {{ font-size: 16pt; margin: 2mm 0 5mm; }}
  h3 {{ font-size: 11pt; margin: 7mm 0 2mm; }}
  table {{ width: 100%; border-collapse: collapse; margin-bottom: 4mm; }}
  th, td {{ border: 1px solid #b8c0c8; padding: 1.6mm 2.4mm; vertical-align: top; text-align: {start}; }}
  .details th {{ width: 32%; background: #f1f3f5; }}
  .lines thead th {{ background: #dde3ea; }}
  .num {{ text-align: {end}; white-space: nowrap; }}
  .grand td {{ font-weight: bold; background: #f1f3f5; }}
</style>
</head>
<body>
{logo_html}
<p class="company">{_e(company_name)}</p>
<h1>{_e(labels['title'])} — {_e(context.get('project_no', ''))}</h1>
<table class="details">{detail_rows}</table>
{body}
</body>
</html>"""
    # Same safety note as cost_workout_service.render_pdf: every value is
    # HTML-escaped and the only URLs are file:// ones built here.
    return HTML(string=html_doc).write_pdf()
